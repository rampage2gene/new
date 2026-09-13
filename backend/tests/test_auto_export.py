"""Every processed document writes its own folder of outputs, and rewrites it after edits."""
from __future__ import annotations

import csv
import json
from pathlib import Path

import pymupdf
import pytest
from openpyxl import load_workbook

from app.exports import auto


@pytest.fixture(scope="module")
def exported(client, fixtures_dir: Path, data_dir: Path) -> tuple[dict, Path]:
    resp = client.post("/api/documents/import", json={"paths": [str(fixtures_dir / "XYZ-5000 Manual (scanned).pdf")]})
    assert resp.status_code == 201, resp.text
    doc = client.get(f"/api/documents/{resp.json()[0]['id']}").json()
    assert doc["status"] == "ready", doc.get("error")
    folder = Path(doc["stats"]["export_dir"])
    return doc, folder


def test_folder_holds_the_six_outputs(exported, data_dir: Path):
    doc, folder = exported
    assert folder.parent == data_dir / "exports"
    names = {p.name for p in folder.iterdir()}
    stem = auto.safe_stem(doc["filename"])
    for suffix in auto.FILES:
        assert f"{stem}.{suffix}" in names, names
    assert set(doc["stats"]["export_files"]) == {f"{stem}.{s}" for s in auto.FILES}
    assert (folder / ".document-id").read_text() == doc["id"]


def test_workbook_has_verified_notes_and_to_fill_sheet(exported):
    doc, folder = exported
    wb = load_workbook(folder / f"{auto.safe_stem(doc['filename'])}.xlsx")
    td = wb["Technical Data"]
    headers = [td.cell(1, c).value for c in range(1, td.max_column + 1)]
    assert "Verified" in headers and "Notes" in headers
    notes_col = headers.index("Notes") + 1
    notes = [td.cell(r, notes_col).value for r in range(2, td.max_row + 1)]
    assert any(n and "confirmed" in n for n in notes)
    assert "To fill in" in wb.sheetnames, "the scanned fixture has values on a single reading"
    tf = wb["To fill in"]
    assert [tf.cell(1, c).value for c in range(1, 5)] == ["Document", "Type", "Page", "Fill in"]
    assert tf.max_row >= 2


def test_clean_pdf_is_text_only_with_a_values_table(exported):
    doc, folder = exported
    pdf = pymupdf.open(folder / f"{auto.safe_stem(doc['filename'])}.clean.pdf")
    text = "".join(p.get_text() for p in pdf)
    assert "Values" in text and "Page 3" in text
    assert "300 A" in text and "Class T" in text
    assert not any(p.get_images() for p in pdf), "no scanned images, only text"


def test_csv_and_json_carry_the_verification(exported):
    doc, folder = exported
    stem = auto.safe_stem(doc["filename"])
    rows = list(csv.DictReader(open(folder / f"{stem}.values.csv", encoding="utf-8")))
    assert rows and {"verified", "verification_status", "notes"} <= set(rows[0])
    assert any(r["verification_status"] == "confirmed" for r in rows)
    payload = json.loads((folder / f"{stem}.json").read_text(encoding="utf-8"))
    assert payload["values"] and payload["verification"]["reader1"]
    assert all("verification" in v for v in payload["values"])


def test_edits_rewrite_the_folder_and_blanks_stay_blank(client, exported):
    doc, folder = exported
    stem = auto.safe_stem(doc["filename"])
    ents = client.get(f"/api/documents/{doc['id']}/entities").json()
    target = next(e for e in ents if e["entity_type"] == "current" and e["ocr_confidence"] is not None)
    other = next(e for e in ents if e["entity_type"] == "voltage" and e["ocr_confidence"] is not None and e["id"] != target["id"])
    assert client.patch(f"/api/entities/{target['id']}", json={"value_text": "127 A"}).status_code == 200
    assert client.patch(f"/api/entities/{other['id']}", json={"value_text": ""}).status_code == 200
    auto.flush_exports()

    rows = list(csv.DictReader(open(folder / f"{stem}.values.csv", encoding="utf-8")))
    filled = next(r for r in rows if r["verification_status"] == "user")
    assert filled["value_text"] == "127 A" and filled["verified"] == "yes"
    blank = next(r for r in rows if r["verification_status"] == "to_fill")
    assert blank["value_text"] == "" and blank["value"] == "" and "TO FILL IN" in blank["notes"]

    wb = load_workbook(folder / f"{stem}.xlsx")
    td = wb["Technical Data"]
    headers = [td.cell(1, c).value for c in range(1, td.max_column + 1)]
    val_col, ver_col = headers.index("Value") + 1, headers.index("Verified") + 1
    cells = [(td.cell(r, val_col).value, td.cell(r, ver_col).value) for r in range(2, td.max_row + 1)]
    assert (127, "yes") in cells
    assert any(v is None for v, _ in cells), "a blank is an empty cell, not a guess"

    pdf = pymupdf.open(folder / f"{stem}.clean.pdf")
    text = "".join(p.get_text() for p in pdf)
    assert "127 A" in text and auto.FILL_MARK in text


def test_open_folder_name_is_stable_and_unique(tmp_path: Path):
    class Doc:
        def __init__(self, id, filename):
            self.id, self.filename = id, filename

    a = auto.export_folder(Doc("a" * 32, "Manual v2.pdf"), tmp_path)
    assert a == tmp_path / "Manual v2"
    a.mkdir()
    (a / ".document-id").write_text("a" * 32)
    assert auto.export_folder(Doc("a" * 32, "Manual v2.pdf"), tmp_path) == a
    b = auto.export_folder(Doc("b" * 32, "Manual v2.pdf"), tmp_path)
    assert b == tmp_path / "Manual v2-bbbbbb"
    assert auto.safe_stem("weird:name*?.PDF") == "weird_name_"
