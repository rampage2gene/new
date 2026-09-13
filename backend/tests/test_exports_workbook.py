"""Workbook export: sheets, live formulas, reference tables, invoice totals."""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _formulas(ws) -> dict[str, str]:
    return {c.coordinate: c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")}


def test_workbook_has_all_sheets_and_links(client, manual_doc, invoice_doc):
    r = client.get("/api/export/workbook")
    assert r.status_code == 200 and r.headers["content-type"].startswith(XLSX)
    wb = load_workbook(io.BytesIO(r.content))
    for name in ("README", "Technical Data", "Reference", "DC Current", "Inverter DC Current", "Voltage Drop", "Battery Runtime", "Alternator Charging", "AC Load", "Fuse & Circuit Protection", "Invoices"):
        assert name in wb.sheetnames, wb.sheetnames
    td = wb["Technical Data"]
    assert td["A1"].value == "Document" and td.max_row > 10
    # numeric values are numbers, and the Source column links back to the page in the app
    values = [td.cell(r, 3).value for r in range(2, td.max_row + 1)]
    assert any(isinstance(v, (int, float)) for v in values)
    links = [td.cell(r, 17).value for r in range(2, td.max_row + 1)]
    assert all(isinstance(l, str) and l.startswith("=HYPERLINK(") and "/documents/" in l and "page=" in l for l in links)
    assert "TechnicalData" in td.tables
    # named ranges the formulas depend on
    for name in ("AwgTable", "Mm2Table", "FuseSizes", "DeviceTable"):
        assert name in wb.defined_names


def test_calculator_sheets_are_live_formulas(client, manual_doc):
    r = client.get("/api/export/workbook", params={"include": "calculators", "calculators": ["fuse_protection", "inverter_dc_current"], "document_ids": [manual_doc["id"]]})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["Fuse & Circuit Protection"]
    f = _formulas(ws)
    assert f, "no formulas on the calculator sheet"
    joined = "\n".join(f.values())
    # results reference input cells (N() around number inputs) and the reference tables, never literals only
    assert "N($B$" in joined and "FuseSizes" in joined and "AwgTable" in joined and "DeviceTable" in joined
    assert "{" not in joined and "template error" not in joined  # every placeholder expanded
    # inputs were prefilled from the manual: continuous current 125 A and the 300 A manufacturer fuse
    labels = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(7, 20)}
    assert labels.get("Maximum continuous current") == 125
    assert labels.get("Manufacturer-specified fuse (optional)") == 300
    # the app's own results sit next to the formulas so the two can be compared after recalculation
    app_col = {ws.cell(r, 1).value: ws.cell(r, 5).value for r in range(1, ws.max_row + 1)}
    assert app_col.get("Manufacturer-specified fuse") == 300
    assert app_col.get("Calculated fuse size (engineering estimate)") == 175
    assert app_col.get("Recommended protection (pending verification)") == 300
    inv = wb["Inverter DC Current"]
    assert any("/" in v and "IF(N(" in v for v in _formulas(inv).values())


def test_single_calculator_export_and_error(client):
    r = client.post("/api/calculators/voltage_drop/export", json={"inputs": {"voltage": 12, "current": 30, "length": 5, "length_unit": "m", "size": "6 AWG", "material": "copper"}})
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('voltage_drop.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb[wb.sheetnames[0]]
    f = _formulas(ws)
    assert any("VLOOKUP" in v and "AwgTable" in v for v in f.values())
    app_col = {(ws.cell(r, 1).value, ws.cell(r, 3).value): ws.cell(r, 5).value for r in range(1, ws.max_row + 1)}
    assert abs(app_col[("Voltage drop", "V")] - 30 * (0.3951 / 304.8) * 10) < 0.01  # 6 AWG Cu, 5 m one-way
    assert abs(app_col[("Voltage drop", "%")] - 3.24) < 0.01
    # a calculation the app cannot run is a 422, not a broken file
    r = client.post("/api/calculators/dc_current/export", json={"inputs": {"power": 100}})
    assert r.status_code == 422


def test_invoice_sheet_formulas(client, invoice_doc):
    r = client.get("/api/export/workbook", params={"include": "invoices", "document_ids": [invoice_doc["id"]]})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content))["Invoices"]
    f = _formulas(ws)
    line_totals = [v for k, v in f.items() if k.startswith("I") and re.fullmatch(r"=F\d+\*H\d+", v)]
    assert len(line_totals) >= 3
    assert any(v.startswith("=SUMIF(") for v in f.values())
    assert any(v.startswith("=SUM(I") for v in f.values())
