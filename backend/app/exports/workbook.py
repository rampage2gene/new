"""Excel workbook builder.

Produces spreadsheets the user can work in: documented values as numeric
cells with a link back to the source page, calculator sheets whose inputs are
editable cells and whose results are real Excel formulas (change an input and
the sheet recalculates), reference tables the formulas look up, and invoice
sheets with SUM/SUMIF totals. openpyxl writes formulas as strings; they are
evaluated by Excel/LibreOffice when the file is opened.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..api.serializers import entity_dict
from ..calculators import tables as T
from ..calculators.base import CalcResult, CalculationError, Calculator, InputValue
from ..calculators.modules import DEVICE_PROFILES, REGISTRY, get_calculator
from ..calculators.suggest import prefill_from_suggestions, suggest_inputs
from ..models import Document, Entity, Invoice

# --------------------------------------------------------------------------- styles
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
INPUT_FILL = PatternFill("solid", fgColor="E8F5E9")  # green: editable input
HELPER_FILL = PatternFill("solid", fgColor="F3F4F6")  # grey: intermediate formula
RESULT_FILL = PatternFill("solid", fgColor="FFF8E1")  # amber: result formula
DOC_FONT = Font(color="1D4ED8")  # blue: value taken from a document
FORMULA_FONT = Font(bold=True)
MUTED_FONT = Font(color="6B7280", italic=True)
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

CLASS_LABEL = {
    "manufacturer_required": "MANUFACTURER REQUIRED",
    "documented_value": "DOCUMENTED VALUE",
    "calculated_estimate": "CALCULATED ESTIMATE",
    "recommended_pending_verification": "RECOMMENDED - PENDING VERIFICATION",
}

_PLACEHOLDER = re.compile(r"\{([a-z_]+(?::[a-z_0-9]+)?)\}")


def _safe_title(name: str) -> str:
    return re.sub(r"[\[\]\*\?/\\:]", "-", name)[:31]


def _header(ws: Worksheet, row: int, labels: list[str], widths: list[int] | None = None) -> None:
    for c, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BOX
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w


def _link(base_url: str | None, document_id: str | None, page: int | None, bbox: list | None, label: str) -> str:
    if not (base_url and document_id and page):
        return label
    url = f"{base_url}/documents/{document_id}?page={page}"
    if bbox and len(bbox) == 4 and all(v is not None for v in bbox):
        url += "&bbox=" + ",".join(f"{float(v):.1f}" for v in bbox)
    return f'=HYPERLINK("{url}","{label}")'


def _num_or_text(v: Any) -> Any:
    """Convert numeric-looking strings to numbers so the sheet can compute on them."""
    if isinstance(v, (int, float)) or v is None:
        return v
    s = str(v).strip().replace(",", "")
    try:
        f = float(s)
        return int(f) if f.is_integer() and "." not in s else f
    except ValueError:
        return v


# --------------------------------------------------------------------------- reference sheet

def write_reference_sheet(wb: Workbook) -> Worksheet:
    ws = wb.create_sheet("Reference")
    ws["A1"] = "Lookup tables used by the calculator formulas. Typical published values - verify against the applicable standard and the actual wire specification."
    ws["A1"].font = MUTED_FONT
    _header(ws, 3, ["AWG", "Ω per 1000 ft (Cu, 20 °C)", "Ampacity 105 °C (A)", "mm²"], [10, 24, 20, 10])
    awg_rows = sorted(T.AWG_OHMS_PER_KFT_CU, key=lambda k: T.AWG_MM2[k], reverse=True)
    for i, key in enumerate(awg_rows, start=4):
        ws.cell(row=i, column=1, value=key)
        ws.cell(row=i, column=2, value=T.AWG_OHMS_PER_KFT_CU[key])
        ws.cell(row=i, column=3, value=T.AWG_AMPACITY_105C.get(key))
        ws.cell(row=i, column=4, value=T.AWG_MM2[key])
    awg_end = 3 + len(awg_rows)

    _header_at(ws, 3, 6, ["mm²", "Ampacity 105 °C (A)"], [10, 20])
    mm2_rows = sorted(T.MM2_AMPACITY_105C)
    for i, key in enumerate(mm2_rows, start=4):
        ws.cell(row=i, column=6, value=key)
        ws.cell(row=i, column=7, value=T.MM2_AMPACITY_105C[key])
    mm2_end = 3 + len(mm2_rows)

    _header_at(ws, 3, 9, ["Standard fuse sizes (A)"], [22])
    for i, s in enumerate(T.STANDARD_FUSE_SIZES, start=4):
        ws.cell(row=i, column=9, value=s)
    fuse_end = 3 + len(T.STANDARD_FUSE_SIZES)

    _header_at(ws, 3, 11, ["Standard breaker sizes (A)"], [24])
    for i, s in enumerate(T.STANDARD_BREAKER_SIZES, start=4):
        ws.cell(row=i, column=11, value=s)

    _header_at(ws, 3, 13, ["Device type key", "Load factor k", "Device", "Fuse characteristic"], [18, 14, 34, 60])
    for i, (key, prof) in enumerate(DEVICE_PROFILES.items(), start=4):
        ws.cell(row=i, column=13, value=key)
        ws.cell(row=i, column=14, value=prof["factor"])
        ws.cell(row=i, column=15, value=prof["label"])
        ws.cell(row=i, column=16, value=prof["char"]).alignment = WRAP
    dev_end = 3 + len(DEVICE_PROFILES)

    ws.cell(row=dev_end + 2, column=13, value="Other constants").font = Font(bold=True)
    for j, (label, value) in enumerate([
        ("Aluminium resistance factor vs copper", T.AL_FACTOR),
        ("Engine-space ampacity derating (105 °C)", T.ENGINE_SPACE_DERATE_105C),
        ("Copper resistivity Ω·mm²/m", T.RHO_CU),
        ("Aluminium resistivity Ω·mm²/m", T.RHO_AL),
    ], start=dev_end + 3):
        ws.cell(row=j, column=13, value=label)
        ws.cell(row=j, column=14, value=value)

    names = {
        "AwgTable": f"Reference!$A$4:$D${awg_end}",
        "Mm2Table": f"Reference!$F$4:$G${mm2_end}",
        "FuseSizes": f"Reference!$I$4:$I${fuse_end}",
        "DeviceTable": f"Reference!$M$4:$P${dev_end}",
    }
    for name, ref in names.items():
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    return ws


def _header_at(ws: Worksheet, row: int, col: int, labels: list[str], widths: list[int] | None = None) -> None:
    for k, label in enumerate(labels):
        cell = ws.cell(row=row, column=col + k, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BOX
        if widths:
            ws.column_dimensions[get_column_letter(col + k)].width = widths[k]


# --------------------------------------------------------------------------- calculator sheet

def _expand(template: str, addr: dict[str, str], kinds: dict[str, str]) -> str:
    def repl(m: re.Match) -> str:
        tag = m.group(1)
        pre, _, key = tag.partition(":")
        if not key:  # plain {input}
            key, pre = pre, ""
        if pre == "":
            cell = addr[f"in:{key}"]
            return f"N({cell})" if kinds.get(key) == "number" else cell
        if pre == "raw":
            return addr[f"in:{key}"]
        if pre == "pct":
            c = addr[f"in:{key}"]
            return f"IF(N({c})>1,N({c})/100,N({c}))"
        if pre == "h":
            return addr[f"h:{key}"]
        if pre == "r":
            return addr[f"r:{key}"]
        raise KeyError(tag)

    return _PLACEHOLDER.sub(repl, template)


def write_calculator_sheet(
    wb: Workbook,
    calc: Calculator,
    payload: dict[str, Any],
    base_url: str | None = None,
    title: str | None = None,
) -> tuple[Worksheet, CalcResult | None, str | None]:
    """Inputs as editable cells, results as live formulas, app-computed values alongside.

    Returns (sheet, python_result_or_None, error_message_or_None).
    """
    spec = calc.spec
    ws = wb.create_sheet(_safe_title(title or spec.name))
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 8
    ws.column_dimensions["D"].width = 34
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["F"].width = 60

    ws["A1"] = spec.name
    ws["A1"].font = TITLE_FONT
    ws["A2"] = spec.formula
    ws["A2"].font = Font(name="Consolas")
    ws["A3"] = spec.description
    ws["A3"].font = MUTED_FONT
    ws["A4"] = "Green cells are inputs: change them and the amber result cells recalculate. Blue values were taken from a document (see Source)."
    ws["A4"].font = MUTED_FONT

    # Python-side computation for the "as computed by app" column
    result: CalcResult | None = None
    error: str | None = None
    try:
        result = calc.run(payload)
    except CalculationError as exc:
        error = str(exc)

    inputs: dict[str, InputValue] = result.inputs if result else {}

    row = 6
    _header(ws, row, ["Input", "Value", "Unit", "Origin", "Source", "Notes"])
    addr: dict[str, str] = {}
    kinds: dict[str, str] = {}
    for inp in spec.inputs:
        row += 1
        kinds[inp.key] = inp.kind
        iv = inputs.get(inp.key)
        raw = payload.get(inp.key)
        if iv is not None:
            value, unit, source, origin = iv.value, iv.unit, iv.source, iv.origin
        else:  # calculation failed before coercion; show what we were given
            value = raw.get("value") if isinstance(raw, dict) else raw
            unit, source, origin = inp.unit, None, "user" if value not in (None, "") else "default"
            if value in (None, ""):
                value = inp.default
        if inp.kind == "number":
            value = _num_or_text(value)
        cell = ws.cell(row=row, column=2, value=value if value is not None else None)
        cell.fill = INPUT_FILL
        cell.border = BOX
        if source:
            cell.font = DOC_FONT
        ws.cell(row=row, column=1, value=inp.label)
        ws.cell(row=row, column=3, value=unit or inp.unit)
        ws.cell(row=row, column=4, value={"document": "from document", "default": "default", "user": "entered"}.get(origin, origin))
        if source and source.page:
            label = f"{source.document_name or 'document'}, p. {source.page}"
            ws.cell(row=row, column=5, value=_link(base_url, source.document_id, source.page, source.bbox, label))
            if source.snippet:
                ws.cell(row=row, column=6, value=source.snippet).alignment = WRAP
        else:
            note = inp.help or ""
            if inp.kind == "select" and inp.options:
                note = (note + " " if note else "") + "Options: " + ", ".join(str(o["value"]) for o in inp.options)
            ws.cell(row=row, column=6, value=note).font = MUTED_FONT
        addr[f"in:{inp.key}"] = f"$B${row}"

    row += 2
    _header(ws, row, ["Calculation", "Value", "Unit", "Classification", "As computed by app", "Notes"])
    app_values = {r.key: r for r in result.results} if result else {}
    for spec_row in spec.excel:
        row += 1
        kind = spec_row.get("kind", "result")
        key = spec_row["key"]
        try:
            formula = _expand(spec_row["formula"], addr, kinds)
        except KeyError as exc:  # template bug; keep the sheet usable
            formula = f"template error: {exc}"
        ws.cell(row=row, column=1, value=spec_row["label"])
        cell = ws.cell(row=row, column=2, value=formula)
        cell.border = BOX
        ws.cell(row=row, column=3, value=spec_row.get("unit"))
        if kind == "helper":
            cell.fill = HELPER_FILL
            ws.cell(row=row, column=4, value="intermediate").font = MUTED_FONT
            addr[f"h:{key}"] = f"$B${row}"
        else:
            cell.fill = RESULT_FILL
            cell.font = FORMULA_FONT
            cls = spec_row.get("classification") or (app_values[key].classification if key in app_values else "calculated_estimate")
            ws.cell(row=row, column=4, value=CLASS_LABEL.get(cls, cls))
            addr[f"r:{key}"] = f"$B${row}"
            if key in app_values:
                rv = app_values[key]
                ws.cell(row=row, column=5, value=rv.value)
                if rv.note:
                    ws.cell(row=row, column=6, value=rv.note).alignment = WRAP
        if kind == "text":
            cell.alignment = WRAP
    if error:
        row += 1
        ws.cell(row=row, column=1, value="App could not compute this calculation with the given inputs:").font = MUTED_FONT
        ws.cell(row=row, column=2, value=error).alignment = WRAP
    if result:
        for rv in result.results:
            if f"r:{rv.key}" not in addr:  # results the sheet has no formula for (e.g. text-only)
                row += 1
                ws.cell(row=row, column=1, value=rv.label)
                ws.cell(row=row, column=2, value=rv.value)
                ws.cell(row=row, column=3, value=rv.unit)
                ws.cell(row=row, column=4, value=CLASS_LABEL.get(rv.classification, rv.classification))
                ws.cell(row=row, column=6, value="computed by the app; no spreadsheet formula").font = MUTED_FONT

    row += 2
    if result and result.steps:
        ws.cell(row=row, column=1, value="Calculation steps (as computed by the app)").font = Font(bold=True)
        for s in result.steps:
            row += 1
            ws.cell(row=row, column=1, value=s).alignment = WRAP
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
    if result and result.warnings:
        ws.cell(row=row, column=1, value="Warnings").font = Font(bold=True, color="B45309")
        for w in result.warnings:
            row += 1
            ws.cell(row=row, column=1, value=w).alignment = WRAP
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
    assumptions = result.assumptions if result else []
    if assumptions:
        ws.cell(row=row, column=1, value="Assumptions").font = Font(bold=True)
        for a in assumptions:
            row += 1
            ws.cell(row=row, column=1, value=a).alignment = WRAP
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        row += 1
    row += 1
    disclaimer = result.disclaimer if result else CalcResult.__dataclass_fields__["disclaimer"].default
    ws.cell(row=row, column=1, value=disclaimer).font = MUTED_FONT
    ws.cell(row=row, column=1).alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
    ws.row_dimensions[row].height = 45
    ws.freeze_panes = "A7"
    return ws, result, error


# --------------------------------------------------------------------------- data sheets

ENTITY_HEADERS = [
    ("Document", 28), ("Type", 14), ("Value", 10), ("Unit", 8), ("As written", 16), ("Qualifier", 12),
    ("Application", 26), ("Circuit", 8), ("Equipment", 16), ("Page", 6), ("Section", 26), ("Confidence", 11),
    ("Flags", 30), ("Snippet", 70), ("Source", 12),
]


def write_technical_data_sheet(wb: Workbook, db: Session, docs: list[Document], base_url: str | None) -> Worksheet:
    ws = wb.create_sheet("Technical Data")
    _header(ws, 1, [h for h, _ in ENTITY_HEADERS], [w for _, w in ENTITY_HEADERS])
    row = 1
    for doc in docs:
        name = doc.title or doc.filename
        ents = db.execute(select(Entity).where(Entity.document_id == doc.id).order_by(Entity.page_number, Entity.char_start)).scalars().all()
        for e in ents:
            d = entity_dict(e, name)
            row += 1
            flags = "; ".join(f.get("message", "") for f in d["flags"])
            values = [
                name, d["entity_type"], d["value"], d["unit"], d["value_text"], d["qualifier"], d["application"], d["circuit"],
                d["equipment"], d["page"], d["section"], d["confidence"], flags, d["snippet"],
                _link(base_url, doc.id, d["page"], d["bbox"], f"p. {d['page']}"),
            ]
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                if c == 3 and v is not None:
                    cell.font = DOC_FONT
            ws.cell(row=row, column=14).alignment = Alignment(wrap_text=False)
    if row > 1:
        table = Table(displayName="TechnicalData", ref=f"A1:{get_column_letter(len(ENTITY_HEADERS))}{row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        ws.add_table(table)
    ws.freeze_panes = "A2"
    return ws


def write_tables_sheet(wb: Workbook, docs: list[Document], base_url: str | None) -> Worksheet | None:
    blocks = []
    for doc in docs:
        for t in (doc.structure or {}).get("tables", []):
            if t.get("rows"):
                blocks.append((doc, t))
    if not blocks:
        return None
    ws = wb.create_sheet("Tables")
    ws["A1"] = "Tables detected in the documents, one block per table. Numeric cells are numbers, so you can reference them in formulas."
    ws["A1"].font = MUTED_FONT
    row = 3
    for doc, t in blocks:
        name = doc.title or doc.filename
        title = f"{name} - page {t['page']}" + (f" - {t['section']}" if t.get("section") else "")
        ws.cell(row=row, column=1, value=title).font = Font(bold=True)
        ws.cell(row=row, column=2, value=_link(base_url, doc.id, t["page"], t.get("bbox"), "open page"))
        row += 1
        for r_i, r in enumerate(t["rows"]):
            for c, v in enumerate(r, start=1):
                cell = ws.cell(row=row, column=c, value=_num_or_text(v))
                cell.border = BOX
                if r_i == 0:
                    cell.font = Font(bold=True)
                    cell.fill = HELPER_FILL
            row += 1
        row += 1
    for c in range(1, 12):
        ws.column_dimensions[get_column_letter(c)].width = 22
    return ws


def write_invoices_sheet(wb: Workbook, db: Session, docs: list[Document], base_url: str | None) -> Worksheet | None:
    doc_by_id = {d.id: d for d in docs}
    invoices = db.execute(select(Invoice).where(Invoice.document_id.in_(list(doc_by_id)))).scalars().all() if docs else []
    invoices = [inv for inv in invoices if inv.line_items]
    if not invoices:
        return None
    ws = wb.create_sheet("Invoices")
    headers = ["Invoice", "Vendor", "Date", "Currency", "Description", "Qty", "Unit", "Unit price", "Line total", "Page"]
    _header(ws, 1, headers, [18, 22, 12, 9, 46, 8, 8, 12, 14, 6])
    row = 1
    first_line = 2
    for inv in invoices:
        doc = doc_by_id[inv.document_id]
        key = inv.invoice_number or (doc.title or doc.filename)
        for li in inv.line_items or []:
            row += 1
            qty = _num_or_text(li.get("quantity"))
            price = _num_or_text(li.get("unit_price"))
            values = [key, inv.vendor, inv.invoice_date, inv.currency, li.get("description"), qty, li.get("unit"), price]
            for c, v in enumerate(values, start=1):
                ws.cell(row=row, column=c, value=v)
            ws.cell(row=row, column=6).fill = INPUT_FILL
            ws.cell(row=row, column=8).fill = INPUT_FILL
            if isinstance(qty, (int, float)) and isinstance(price, (int, float)):
                ws.cell(row=row, column=9, value=f"=F{row}*H{row}").fill = RESULT_FILL
            else:  # keep the printed total when qty/price were not both readable
                ws.cell(row=row, column=9, value=_num_or_text(li.get("total")))
            ws.cell(row=row, column=10, value=_link(base_url, doc.id, li.get("page"), li.get("bbox"), f"p. {li.get('page')}") if li.get("page") else None)
    last_line = row
    table = Table(displayName="InvoiceLines", ref=f"A1:J{last_line}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)

    row += 2
    _header_at(ws, row, 1, ["Totals by invoice", "Vendor", "Lines", "Subtotal (printed)", "Tax (printed)", "Total (printed)", "", "", "Computed from lines"])
    for inv in invoices:
        row += 1
        doc = doc_by_id[inv.document_id]
        key = inv.invoice_number or (doc.title or doc.filename)
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=inv.vendor)
        ws.cell(row=row, column=3, value=f'=COUNTIF($A${first_line}:$A${last_line},A{row})')
        ws.cell(row=row, column=4, value=inv.subtotal)
        ws.cell(row=row, column=5, value=inv.tax)
        ws.cell(row=row, column=6, value=inv.total)
        ws.cell(row=row, column=9, value=f'=SUMIF($A${first_line}:$A${last_line},A{row},$I${first_line}:$I${last_line})').fill = RESULT_FILL
    row += 1
    ws.cell(row=row, column=1, value="JOB TOTAL").font = Font(bold=True)
    ws.cell(row=row, column=6, value=f"=SUM(F{row - len(invoices)}:F{row - 1})").font = FORMULA_FONT
    ws.cell(row=row, column=9, value=f"=SUM(I{first_line}:I{last_line})").font = FORMULA_FONT
    ws.cell(row=row, column=9).fill = RESULT_FILL

    row += 2
    _header_at(ws, row, 1, ["Estimate by item", "", "", "", "Description", "Qty", "", "Avg unit price", "Total"])
    seen: list[str] = []
    for inv in invoices:
        for li in inv.line_items or []:
            d = (li.get("description") or "").strip()
            if d and d.lower() not in [s.lower() for s in seen]:
                seen.append(d)
    est_first = row + 1
    for d in seen:
        row += 1
        ws.cell(row=row, column=5, value=d)
        ws.cell(row=row, column=6, value=f'=SUMIF($E${first_line}:$E${last_line},E{row},$F${first_line}:$F${last_line})')
        ws.cell(row=row, column=8, value=f'=IFERROR(I{row}/F{row},"")')
        ws.cell(row=row, column=9, value=f'=SUMIF($E${first_line}:$E${last_line},E{row},$I${first_line}:$I${last_line})').fill = RESULT_FILL
    row += 1
    ws.cell(row=row, column=5, value="PROJECT TOTAL").font = Font(bold=True)
    ws.cell(row=row, column=9, value=f"=SUM(I{est_first}:I{row - 1})").font = FORMULA_FONT
    ws.freeze_panes = "A2"
    return ws


def write_readme_sheet(wb: Workbook, docs: list[Document], sheets: list[str], note: str | None = None) -> Worksheet:
    ws = wb.active
    ws.title = "README"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    ws["A1"] = "Marine Electrical Document Intelligence - workbook export"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    ws["A2"].font = MUTED_FONT
    row = 4
    ws.cell(row=row, column=1, value="Documents").font = Font(bold=True)
    for d in docs:
        row += 1
        ws.cell(row=row, column=1, value=d.title or d.filename)
        ws.cell(row=row, column=2, value=f"{d.filename} - {d.manufacturer or ''} {d.model_number or ''} - {d.page_count} pages, {d.ocr_pages} OCR'd".strip())
    row += 2
    ws.cell(row=row, column=1, value="Sheets").font = Font(bold=True)
    for s in sheets:
        row += 1
        ws.cell(row=row, column=1, value=s)
    row += 2
    ws.cell(row=row, column=1, value="Colour legend").font = Font(bold=True)
    for label, fill, font in [
        ("Editable input - change it and dependent formulas recalculate", INPUT_FILL, None),
        ("Value taken from a document (blue text); the Source column links to the page", None, DOC_FONT),
        ("Intermediate formula", HELPER_FILL, None),
        ("Result formula", RESULT_FILL, FORMULA_FONT),
    ]:
        row += 1
        c = ws.cell(row=row, column=1, value="")
        if fill:
            c.fill = fill
        ws.cell(row=row, column=2, value=label).font = font or Font()
    row += 2
    ws.cell(row=row, column=1, value="Read this").font = Font(bold=True)
    ws.cell(row=row, column=2, value=CalcResult.__dataclass_fields__["disclaimer"].default).alignment = WRAP
    ws.row_dimensions[row].height = 45
    row += 1
    ws.cell(row=row, column=2, value="Formula results appear when the file is opened in Excel, LibreOffice or Google Sheets. 'As computed by app' columns show the app's own result for the same inputs so the two can be compared.").alignment = WRAP
    if note:
        row += 1
        ws.cell(row=row, column=2, value=note).alignment = WRAP
    return ws


# --------------------------------------------------------------------------- entry points

def build_workbook(
    db: Session,
    document_ids: list[str] | None,
    include: set[str] | None = None,
    base_url: str | None = None,
    calculators: list[str] | None = None,
) -> bytes:
    include = include or {"data", "tables", "calculators", "invoices"}
    stmt = select(Document).where(Document.status == "ready")
    if document_ids:
        stmt = stmt.where(Document.id.in_(document_ids))
    docs = list(db.execute(stmt.order_by(Document.uploaded_at)).scalars().all())
    wb = Workbook()
    sheets: list[str] = []
    notes: list[str] = []
    if "data" in include:
        write_technical_data_sheet(wb, db, docs, base_url)
        sheets.append("Technical Data - every extracted value with page link")
    if "tables" in include and write_tables_sheet(wb, docs, base_url):
        sheets.append("Tables - tables found in the documents")
    if "calculators" in include:
        write_reference_sheet(wb)
        sheets.append("Reference - lookup tables used by the formulas")
        # Prefill each calculator from the document whose values cover the most inputs
        # (ties go to the earliest upload); the sheet stays usable when nothing matches.
        ents_by_doc = {d.id: db.execute(select(Entity).where(Entity.document_id == d.id)).scalars().all() for d in docs}
        for calc_id in calculators or list(REGISTRY):
            calc = get_calculator(calc_id)
            payload: dict[str, Any] = {}
            for d in docs:
                candidate = prefill_from_suggestions(calc, suggest_inputs(calc, ents_by_doc[d.id], d.title or d.filename))
                if len(candidate) > len(payload):
                    payload = candidate
            _, _, err = write_calculator_sheet(wb, calc, payload, base_url)
            sheets.append(f"{calc.spec.name} - inputs prefilled from the documents where possible; results are live formulas")
            if err:
                notes.append(f"{calc.spec.name}: {err} (fill in the green cells)")
    if "invoices" in include and write_invoices_sheet(wb, db, docs, base_url):
        sheets.append("Invoices - line items with computed totals and an estimate by item")
    write_readme_sheet(wb, docs, sheets, "\n".join(notes) if notes else None)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_calculator_workbook(calc_id: str, payload: dict[str, Any], base_url: str | None = None) -> bytes:
    calc = get_calculator(calc_id)
    wb = Workbook()
    write_reference_sheet(wb)
    _, result, err = write_calculator_sheet(wb, calc, payload, base_url)
    if err and not result:
        raise CalculationError(err)
    write_readme_sheet(wb, [], [f"{calc.spec.name} - inputs and live-formula results", "Reference - lookup tables"])
    wb.move_sheet(_safe_title(calc.spec.name), offset=-(len(wb.sheetnames) - 1))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
